# G5：合规报告模板引擎测试（docx 占位符注入）
# 运行：python -m pytest tests/test_report_template.py -q
import os
import sys
import tempfile

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

import pytest

from common.models import SEVERITY_HIGH, SEVERITY_MEDIUM, STATUS_CONFIRMED, ScanResult
from core.report import ReportBuilder
from lib.report_template import build_scalar_values, render_docx_template

try:
    from docx import Document
except ImportError:
    pytest.skip("python-docx 未安装", allow_module_level=True)


def _sample_builder():
    results = [
        ScanResult(
            kind="vuln",
            name="SQL注入漏洞",
            severity=SEVERITY_HIGH,
            status=STATUS_CONFIRMED,
            url="http://x.com/system/role/list",
            evidence="XPATH syntax error",
            fix="使用预编译语句",
        ),
        ScanResult(
            kind="vuln",
            name="任意文件读取",
            severity=SEVERITY_MEDIUM,
            status=STATUS_CONFIRMED,
            url="http://x.com/common/download/resource",
            evidence="root:x:0:0",
            fix="限制下载路径",
        ),
    ]
    builder = ReportBuilder(
        results=results, target="http://x.com", summary={"mode": "u", "duration": 12.5, "request_count": 88}
    )
    return builder


def _make_template(tmpdir, text_placeholder=""):
    """构造安服风格模板：抬头 + 标量占位符 + 表格占位符"""
    doc = Document()
    doc.add_heading("渗透测试报告", level=0)
    doc.add_paragraph("测试目标：{{target}}")
    doc.add_paragraph("扫描日期：{{scan_date}}　确认漏洞：{{confirmed}}（高 {{high}} / 中 {{medium}}）")
    doc.add_paragraph("{{vuln_table}}")
    if text_placeholder:
        doc.add_paragraph(text_placeholder)
    path = os.path.join(str(tmpdir), "tpl.docx")
    doc.save(path)
    return path


def _read_all_text(docx_path):
    doc = Document(docx_path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def test_build_scalar_values():
    builder = _sample_builder()
    vals = build_scalar_values(builder)
    assert vals["target"] == "http://x.com"
    assert vals["total"] == "2"
    assert vals["high"] == "1" and vals["medium"] == "1"
    assert vals["mode"] == "u"
    assert vals["request_count"] == "88"
    assert vals["duration"] == "12.50"


def test_render_docx_template_scalars_and_table(tmp_path):
    """标量替换 + {{vuln_table}} 定点插入表格（含表头与两行明细）"""
    tpl = _make_template(tmp_path)
    out = os.path.join(str(tmp_path), "out.docx")
    result = render_docx_template(tpl, _sample_builder(), out)
    assert result == out

    text = _read_all_text(out)
    # 标量已替换（无残留占位符）
    assert "{{" not in text
    assert "测试目标：http://x.com" in text
    # 表格已插入：表头 + 2 行漏洞
    doc = Document(out)
    tables = doc.tables
    assert len(tables) == 1
    t = tables[0]
    assert t.rows[0].cells[0].text == "漏洞名称"
    assert "SQL注入漏洞" in t.rows[1].cells[0].text
    assert "任意文件读取" in t.rows[2].cells[0].text


def test_render_docx_template_details_block(tmp_path):
    """{{vuln_details}} 逐漏洞详述插入"""
    doc = Document()
    doc.add_paragraph("{{vuln_details}}")
    tpl = os.path.join(str(tmp_path), "tpl2.docx")
    doc.save(tpl)
    out = os.path.join(str(tmp_path), "out2.docx")
    render_docx_template(tpl, _sample_builder(), out)

    text = _read_all_text(out)
    assert "1. SQL注入漏洞（high）" in text
    assert "2. 任意文件读取（medium）" in text
    assert "证据：XPATH syntax error" in text


def test_render_docx_template_empty_confirmed(tmp_path):
    """零确认漏洞：表格与详述给出占位文案而非报错"""
    builder = ReportBuilder(results=[], target="http://empty.com")
    doc = Document()
    doc.add_paragraph("{{vuln_table}}")
    doc.add_paragraph("{{vuln_details}}")
    tpl = os.path.join(str(tmp_path), "tpl3.docx")
    doc.save(tpl)
    out = os.path.join(str(tmp_path), "out3.docx")
    render_docx_template(tpl, builder, out)
    text = _read_all_text(out)
    assert "未发现已确认漏洞" in text


def test_render_docx_template_missing_file(tmp_path):
    """模板不存在 → None（调用方降级默认报告，不抛异常）"""
    out = os.path.join(str(tmp_path), "out4.docx")
    assert render_docx_template(str(tmp_path) + "/nope.docx", _sample_builder(), out) is None


def test_render_docx_template_fail_on_unresolved(tmp_path):
    """未识别占位符 + fail_on_unresolved=True → ValueError（模板校验模式）"""
    tpl = _make_template(tmp_path, text_placeholder="未定义变量：{{oops_unknown}}")
    out = os.path.join(str(tmp_path), "out5.docx")
    with pytest.raises(ValueError):
        render_docx_template(tpl, _sample_builder(), out, fail_on_unresolved=True)


if __name__ == "__main__":
    import pytest

    pytest.main([__file__, "-v", "--tb=short"])
