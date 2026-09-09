# G5：整改复测工作流测试（基线对比 → 整改验证报告 docx/JSON）
# 运行：python -m pytest tests/test_remediation.py -q
import json
import os
import sys
import tempfile

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

import pytest

from common.models import (
    SEVERITY_HIGH,
    SEVERITY_MEDIUM,
    STATUS_CONFIRMED,
    STATUS_SAFE,
    ScanResult,
)
from core.report import ReportBuilder
from lib.remediation import (
    build_remediation_report,
    render_remediation_docx,
    run_remediation,
)

try:
    from docx import Document
except ImportError:
    pytest.skip("python-docx 未安装", allow_module_level=True)


def _result(name, url, severity=SEVERITY_HIGH, status=STATUS_CONFIRMED):
    return ScanResult(
        kind="vuln",
        name=name,
        severity=severity,
        status=status,
        url=url,
        evidence="evidence",
        fix="fix it",
    )


def _report(results, scan_time="2026-09-09 10:00:00"):
    b = ReportBuilder(results=results, target="http://x.com")
    d = b.to_dict()
    d["scan_time"] = scan_time
    return d


def test_build_remediation_classification():
    """分类：closed（整改闭环）/ open（仍未整改）/ new（复测新发现）"""
    old = _report(
        [
            _result("SQL注入", "http://x.com/a"),  # 整改后消失 → closed
            _result("任意文件读取", "http://x.com/b"),  # 复测仍在 → open
        ]
    )
    new = _report(
        [
            _result("任意文件读取", "http://x.com/b"),  # 未整改
            _result("Druid未授权", "http://x.com/c"),  # 复测新发现
        ]
    )
    rep = build_remediation_report(old, new)
    closed_names = {e.name for e in rep.closed}
    open_names = {e.name for e in rep.open_items}
    new_names = {e.name for e in rep.new_findings}
    assert closed_names == {"SQL注入"}
    assert open_names == {"任意文件读取"}
    assert new_names == {"Druid未授权"}
    # 1 closed / (1 closed + 1 open) = 50%
    assert rep.remediation_rate == 50.0
    assert "部分整改" in rep.conclusion


def test_build_remediation_full_closed():
    """全部整改闭环 → 完成率 100% + 闭环结论"""
    old = _report([_result("SQL注入", "http://x.com/a")])
    new = _report([])
    rep = build_remediation_report(old, new)
    assert rep.remediation_rate == 100.0
    assert "整改完成" in rep.conclusion


def test_build_remediation_no_remediation():
    """零整改 → 完成率 0 + 未见整改结论"""
    old = _report([_result("SQL注入", "http://x.com/a")])
    new = _report([_result("SQL注入", "http://x.com/a")])
    rep = build_remediation_report(old, new)
    assert rep.total_closed == 0 and rep.total_open == 1
    assert rep.remediation_rate == 0.0
    assert "未见整改" in rep.conclusion


def test_build_remediation_empty_baseline():
    """基线为零漏洞 → 完成率 100%（无可整改项视为闭环）"""
    rep = build_remediation_report(_report([]), _report([]))
    assert rep.remediation_rate == 100.0


def test_render_remediation_docx(tmp_path):
    """docx 交付物：摘要 + 三个分类表"""
    old = _report(
        [
            _result("SQL注入", "http://x.com/a", SEVERITY_HIGH),
            _result("任意文件读取", "http://x.com/b", SEVERITY_MEDIUM),
        ],
        scan_time="2026-09-08 09:00:00",
    )
    new = _report([_result("任意文件读取", "http://x.com/b", SEVERITY_MEDIUM)])
    rep = build_remediation_report(old, new)
    out = os.path.join(str(tmp_path), "remediation.docx")
    result = render_remediation_docx(rep, out)
    assert result == out

    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "漏洞整改验证报告" in text
    assert "整改完成率：50.0%" in text
    assert "部分整改" in text

    tables_text = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "SQL注入" in tables_text  # closed 表
    assert "任意文件读取" in tables_text  # open 表


def test_run_remediation_outputs(tmp_path):
    """CLI 入口：基线 JSON → remediation.json + remediation.docx"""
    old = _report([_result("SQL注入", "http://x.com/a")])
    baseline_path = os.path.join(str(tmp_path), "baseline.json")
    with open(baseline_path, "w", encoding="utf-8") as f:
        json.dump(old, f, ensure_ascii=False)

    builder = ReportBuilder(
        results=[_result("SQL注入", "http://x.com/a", SEVERITY_HIGH, STATUS_CONFIRMED)], target="http://x.com"
    )
    args = type("Args", (), {"remediation": baseline_path, "report": str(tmp_path)})()
    outputs = run_remediation(args, builder)
    names = [os.path.basename(p) for p in outputs]
    assert names == ["remediation.json", "remediation.docx"]
    with open(outputs[0], encoding="utf-8") as f:
        data = json.load(f)
    assert data["summary"]["closed"] == 0 and data["summary"]["open"] == 1


def test_run_remediation_missing_baseline(tmp_path):
    """基线文件不存在 → 返回空列表不抛异常"""
    builder = ReportBuilder(results=[], target="http://x.com")
    args = type("Args", (), {"remediation": str(tmp_path) + "/nope.json", "report": str(tmp_path)})()
    assert run_remediation(args, builder) == []


if __name__ == "__main__":
    import pytest

    pytest.main([__file__, "-v", "--tb=short"])
