# G5：整改复测工作流（安服第二高频交付物）
#
# 场景：首次扫描发现漏洞（基线）→ 客户整改 → 复测扫描 → 出「整改验证报告」：
#   哪些漏洞已整改闭环（CLOSED）、哪些仍未整改（OPEN）、复测新发现哪些（NEW），
#   并给出整改完成率与结论——安服复测交付的标准化。
#
# 使用方式：
#   # 首次扫描保存基线
#   python main.py -p http://target/ --report reports/ --save-baseline
#   # （客户整改后）复测并出整改验证报告
#   python main.py -p http://target/ --report reports/ --remediation reports/baseline.json
#
# 数据来源：复用 D20 diff_reports 的指纹对比（name|url，仅统计 CONFIRMED），
# 在其分类之上映射安服语义：
#   diff fixed（旧有新无）          → CLOSED 已整改闭环
#   diff persisted / changed        → OPEN  仍未整改（含状态变化的回归项）
#   diff new（新有旧无）            → NEW   复测新发现
from dataclasses import dataclass, field
from typing import Any, Dict, List

from lib.diff_scan import DiffEntry, diff_reports, load_report


@dataclass
class RemediationReport:
    """整改验证报告"""

    target: str = ""
    baseline_scan_time: str = ""
    retest_scan_time: str = ""
    closed: List[DiffEntry] = field(default_factory=list)  # 已整改闭环
    open_items: List[DiffEntry] = field(default_factory=list)  # 仍未整改
    new_findings: List[DiffEntry] = field(default_factory=list)  # 复测新发现

    @property
    def total_closed(self) -> int:
        return len(self.closed)

    @property
    def total_open(self) -> int:
        return len(self.open_items)

    @property
    def total_new(self) -> int:
        return len(self.new_findings)

    @property
    def remediation_rate(self) -> float:
        """整改完成率 = 已闭环 / (已闭环 + 未整改)；基线为零漏洞时视为 100%"""
        denominator = self.total_closed + self.total_open
        if denominator == 0:
            return 100.0
        return round(self.total_closed / denominator * 100.0, 1)

    @property
    def conclusion(self) -> str:
        """整改结论（安服报告结论段用语）"""
        if self.total_open == 0 and self.total_new == 0:
            return "整改完成，全部漏洞已闭环，复测无新发现"
        if self.total_closed > 0:
            return f"部分整改：{self.total_closed} 项已闭环，{self.total_open} 项仍未整改" + (
                f"，复测新发现 {self.total_new} 项" if self.total_new else ""
            )
        return "未见整改：基线漏洞在复测中全部复现" + (f"，且复测新发现 {self.total_new} 项" if self.total_new else "")

    def to_dict(self) -> Dict[str, Any]:
        return {
            "target": self.target,
            "baseline_scan_time": self.baseline_scan_time,
            "retest_scan_time": self.retest_scan_time,
            "summary": {
                "closed": self.total_closed,
                "open": self.total_open,
                "new": self.total_new,
                "remediation_rate": self.remediation_rate,
                "conclusion": self.conclusion,
            },
            "closed": [e.__dict__ for e in self.closed],
            "open_items": [e.__dict__ for e in self.open_items],
            "new_findings": [e.__dict__ for e in self.new_findings],
        }


def build_remediation_report(old_report: Dict[str, Any], new_report: Dict[str, Any]) -> RemediationReport:
    """从两次扫描报告构建整改验证报告

    Args:
        old_report: 整改前基线报告（builder.to_dict() 的 JSON 结构）
        new_report: 复测扫描报告

    Returns:
        RemediationReport（closed/open/new 分类 + 完成率 + 结论）
    """
    diff = diff_reports(old_report, new_report)
    rep = RemediationReport(
        target=diff.target,
        baseline_scan_time=diff.old_scan_time,
        retest_scan_time=diff.new_scan_time,
        closed=list(diff.fixed_vulns),
        new_findings=list(diff.new_vulns),
    )
    # persisted（原样复现）与 changed（状态/严重度变化）都说明"基线漏洞仍存在"
    rep.open_items = list(diff.persisted_vulns) + list(diff.changed_vulns)
    return rep


def render_remediation_docx(rep: RemediationReport, out_path: str) -> str:
    """渲染整改验证报告 docx（安服交付物版式：摘要 → 三个分类表）"""
    from docx import Document
    from docx.shared import RGBColor

    doc = Document()
    doc.add_heading("漏洞整改验证报告", level=0)
    doc.add_paragraph(f"测试目标：{rep.target}")
    doc.add_paragraph(f"基线扫描：{rep.baseline_scan_time or '-'}　复测时间：{rep.retest_scan_time or '-'}")
    doc.add_paragraph(
        f"整改完成率：{rep.remediation_rate}%　"
        f"已闭环 {rep.total_closed} 项 / 未整改 {rep.total_open} 项 / 复测新发现 {rep.total_new} 项"
    )
    doc.add_paragraph(f"结论：{rep.conclusion}")

    def _add_table(title: str, entries: List[DiffEntry], empty_hint: str):
        doc.add_heading(title, level=2)
        if not entries:
            doc.add_paragraph(empty_hint)
            return
        table = doc.add_table(rows=1, cols=4)
        try:
            table.style = "Table Grid"
        except Exception:
            pass
        header = ["漏洞名称", "严重度", "URL", "备注"]
        for j, h in enumerate(header):
            cell = table.rows[0].cells[j]
            cell.text = h
        for e in entries:
            cells = table.add_row().cells
            cells[0].text = e.name
            cells[1].text = e.new_severity or e.old_severity or "-"
            cells[2].text = e.url or ""
            cells[3].text = "整改前状态: %s" % e.old_status if e.old_status else "-"
        # 表头蓝底白字（底色经 tcPr 直接设置，与主报告一致）
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        for cell in table.rows[0].cells:
            for run in cell.paragraphs[0].runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "0969DA")
            cell._tc.get_or_add_tcPr().append(shd)

    _add_table("已整改闭环（CLOSED）", rep.closed, "无（基线漏洞均已整改闭环）")
    _add_table("仍未整改（OPEN）", rep.open_items, "无未整改项")
    _add_table("复测新发现（NEW）", rep.new_findings, "复测无新发现")

    import os

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    doc.save(out_path)
    return out_path


def run_remediation(args, builder: Any) -> List[str]:
    """CLI 入口：--remediation <基线.json>（复测扫描完成后调用）

    Args:
        args: CLI 参数（remediation / report / debug）
        builder: 本次（复测）扫描的 ReportBuilder

    Returns:
        生成的文件路径列表（JSON + docx）；基线加载失败返回 []
    """
    import json
    import os

    outputs: List[str] = []
    try:
        old_report = load_report(args.remediation)
    except (OSError, ValueError) as e:
        from common.logger import get_logger

        get_logger(__name__).warning("整改基线加载失败: %s", e)
        return []
    rep = build_remediation_report(old_report, builder.to_dict())

    out_dir = getattr(args, "report", None) or "reports"
    json_path = os.path.join(out_dir, "remediation.json")
    os.makedirs(out_dir, exist_ok=True)
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(rep.to_dict(), f, ensure_ascii=False, indent=2)
    outputs.append(json_path)

    try:
        docx_path = os.path.join(out_dir, "remediation.docx")
        render_remediation_docx(rep, docx_path)
        outputs.append(docx_path)
    except ImportError:
        pass  # python-docx 未安装，JSON 交付物仍然可用

    return outputs
